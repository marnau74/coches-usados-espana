"""Datos sintéticos con el formato exacto de las fuentes, para los tests y la CI.

Las fuentes reales pesan 2 GB: aquí se fabrican ficheros pequeños y deterministas con el
mismo formato (líneas de 714 caracteres en Latin-1, ZIP con un único .txt, Excel de la DGT,
JSON-stat de Eurostat), y las tablas «oficiales» salen de contar lo que se ha fabricado, así
que los modelos y los tests de cuadre de dbt se pueden ejecutar de punta a punta.

Cada mes de 2024 lleva, por tipo de trámite, un número conocido de filas (ver `PLAN`). Algunas
tienen trampa a propósito: una fila de turismo con fecha del mes anterior (trámite tardío),
una baja temporal, una rematriculación y un municipio con eñe.
"""

from __future__ import annotations

import json
import zipfile
from datetime import date
from pathlib import Path

import openpyxl

from coches_usados.ingest.diseno_registro import CAMPOS, CODIFICACION, LONGITUD_LINEA

ANIO = 2024
MESES = range(1, 13)
COLUMNAS_TABLA = [
    "Camiones", "Furgonetas", "Autobuses", "Turismos", "Motocicletas",
    "Tractores\nindustriales", "Remolques y\nsemirremolques", "Otros\nvehículos",
]  # fmt: skip

# Filas por mes y tipo de vehículo en transferencias (el mes `m` suma `m` a los turismos).
PLAN_TRANSFERENCIAS = {"40": 30, "7A": 2, "20": 10, "02": 5, "50": 8, "81": 1}
GRUPO_DE_TIPO = {
    "40": "Turismos", "7A": "Turismos", "20": "Furgonetas", "02": "Camiones", "50": "Motocicletas", "81": "Tractores\nindustriales",
}  # fmt: skip
PROVINCIAS = ["MU", "M", "B", "V", "SE"]
MUNICIPIOS = ["MURCIA", "MADRID", "BARCELONA", "A CORUÑA", "SEVILLA"]
MARCAS_MODELOS = [
    ("SEAT", "IBIZA 1.0 TSI"),
    ("NISSAN", "NISSAN QASHQAI"),
    ("MERCEDES", "A 200 D"),
    ("BMW", "320D"),
    ("CITROEN", "NUEVO C3"),
    ("VOLKSWAGEN", "GOLF 1.9 TDI"),
]
# (código de propulsión, categoría eléctrica)
COMBUSTIBLES = [("0", ""), ("1", ""), ("0", "HEV"), ("2", "BEV"), ("1", "PHEV"), ("4", "")]


def _linea(**campos: str) -> str:
    """Una línea de 714 caracteres: cada campo en su posición, relleno con espacios."""
    nombres = {c.nombre for c in CAMPOS}
    desconocidos = set(campos) - nombres
    if desconocidos:
        raise KeyError(f"Campos que no existen en el diseño de registro: {sorted(desconocidos)}")
    partes = []
    for c in CAMPOS:
        valor = str(campos.get(c.nombre, ""))
        if len(valor) > c.longitud:
            raise ValueError(f"{c.nombre}: «{valor}» no cabe en {c.longitud} caracteres")
        partes.append(valor.ljust(c.longitud))
    linea = "".join(partes)
    assert len(linea) == LONGITUD_LINEA
    return linea


def _fecha(anio: int, mes: int, dia: int) -> str:
    return f"{dia:02d}{mes:02d}{anio}"


def _fila(i: int, tramite_fecha: str, clave: str, cod_tipo: str, **extra: str) -> str:
    marca, modelo = MARCAS_MODELOS[i % len(MARCAS_MODELOS)]
    propulsion, categoria = COMBUSTIBLES[i % len(COMBUSTIBLES)]
    provincia = PROVINCIAS[i % len(PROVINCIAS)]
    anio_matricula = 2005 + (i * 7) % 19  # de 2005 a 2023, siempre anterior al trámite
    datos = {
        "fec_matricula": _fecha(anio_matricula, 1 + i % 12, 1 + i % 28),
        "fec_tramitacion": tramite_fecha,
        "marca_itv": marca,
        "modelo_itv": modelo,
        "cod_tipo": cod_tipo,
        "cod_propulsion_itv": propulsion,
        "categoria_vehiculo_electrico": categoria,
        "localidad_vehiculo": MUNICIPIOS[i % len(MUNICIPIOS)],
        "cod_provincia_veh": provincia,
        "cod_provincia_mat": provincia,
        "clave_tramite": clave,
        "fec_tramite": tramite_fecha,
        "ind_nuevo_usado": "N" if i % 5 else "U",
        "persona_fisica_juridica": "D" if i % 4 else "X",
        "municipio": MUNICIPIOS[i % len(MUNICIPIOS)],
        "fec_proceso": tramite_fecha,
    }
    datos.update(extra)
    return _linea(**datos)


def lineas_transferencias(mes: int) -> list[str]:
    lineas, i = [], 0
    for cod_tipo, base in PLAN_TRANSFERENCIAS.items():
        cuantas = base + (mes if cod_tipo == "40" else 0)
        for _ in range(cuantas):
            lineas.append(_fila(i, _fecha(ANIO, mes, 1 + i % 28), "2", cod_tipo))
            i += 1
    # Trámite tardío: un turismo fechado el mes anterior (en enero, anterior a la ventana).
    anterior = (ANIO, mes - 1) if mes > 1 else (ANIO - 1, 12)
    lineas.append(_fila(i, _fecha(*anterior, 15), "2", "40"))
    return lineas


def lineas_matriculaciones(mes: int) -> list[str]:
    lineas = [_fila(i, _fecha(ANIO, mes, 1 + i % 28), "1", "40") for i in range(20)]
    lineas.append(_fila(21, _fecha(ANIO, mes, 5), "5", "40"))  # rematriculación: no cuenta
    lineas.append(_fila(22, _fecha(ANIO, mes, 6), "1", "50"))  # moto: no es turismo
    return lineas


def lineas_bajas(mes: int) -> list[str]:
    lineas, i = [], 0
    for clave, motivo, cuantas in (("3", "7", 12), ("3", "4", 2), ("7", "8", 3), ("6", "", 5), ("3", "A", 1)):
        for _ in range(cuantas):
            lineas.append(_fila(i, _fecha(ANIO, mes, 1 + i % 28), clave, "40", ind_baja_def=motivo))
            i += 1
    lineas.append(_fila(i, _fecha(ANIO, mes, 3), "3", "20", ind_baja_def="7"))  # furgoneta
    return lineas


def _escribir_zip(destino: Path, nombre: str, lineas: list[str]) -> None:
    destino.parent.mkdir(parents=True, exist_ok=True)
    texto = "\r\n".join(lineas) + "\r\n"
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"{nombre}.txt", texto.encode(CODIFICACION))


def escribir_microdatos(raw: Path) -> None:
    generadores = {
        "transferencias": ("trf", lineas_transferencias),
        "matriculaciones": ("mat", lineas_matriculaciones),
        "bajas": ("bajas", lineas_bajas),
    }
    for carpeta, (clave, generar) in generadores.items():
        for mes in MESES:
            nombre = f"export_mensual_{clave}_{ANIO}{mes:02d}"
            _escribir_zip(raw / "dgt" / carpeta / f"{nombre}.zip", nombre, generar(mes))


def _recuento(lineas: list[str], por_columna) -> dict[tuple[int, str], int]:
    """Filas por (mes de la fecha de trámite, columna de la tabla de la DGT)."""
    cuenta: dict[tuple[int, str], int] = {}
    campos = {c.nombre: c for c in CAMPOS}
    for linea in lineas:
        f = campos["fec_tramite"]
        fecha = linea[f.inicio - 1 : f.inicio - 1 + f.longitud]
        anio, mes = int(fecha[4:]), int(fecha[2:4])
        if anio != ANIO:
            continue
        columna = por_columna(linea, campos)
        if columna:
            cuenta[(mes, columna)] = cuenta.get((mes, columna), 0) + 1
    return cuenta


def _valor(linea: str, campos: dict, nombre: str) -> str:
    c = campos[nombre]
    return linea[c.inicio - 1 : c.inicio - 1 + c.longitud].strip()


def _hoja_mensual(
    wb: openpyxl.Workbook, nombre_hoja: str, titulo: str, cuenta: dict[tuple[int, str], int], con_totales: bool
) -> None:
    hoja = wb.create_sheet(nombre_hoja)
    columnas = COLUMNAS_TABLA if not con_totales else [*COLUMNAS_TABLA[:7], COLUMNAS_TABLA[7]]
    hoja.append([titulo])
    hoja.append([])
    hoja.append(["Mes", *columnas, "Total"] + (["Total general"] if con_totales else []))
    nombres = [
        "Enero",
        "Febrero",
        "Marzo",
        "Abril",
        "Mayo",
        "Junio",
        "Julio",
        "Agosto",
        "Septiembre",
        "Octubre",
        "Noviembre",
        "Diciembre",
    ]
    totales = [0] * len(columnas)
    for mes in MESES:
        fila = [cuenta.get((mes, c), 0) for c in columnas]
        totales = [t + v for t, v in zip(totales, fila, strict=True)]
        hoja.append([nombres[mes - 1], *fila, sum(fila)] + ([sum(fila)] if con_totales else []))
    hoja.append(["Total", *totales, sum(totales)] + ([sum(totales)] if con_totales else []))


def escribir_tablas_dgt(raw: Path) -> dict[str, int]:
    """Excel de la DGT: tablas mensuales de transferencias y bajas y serie histórica. Las cifras
    salen de contar las filas fabricadas (no de la lógica de los modelos)."""
    carpeta = raw / "dgt_tablas"
    carpeta.mkdir(parents=True, exist_ok=True)

    def grupo_transferencia(linea: str, campos: dict) -> str | None:
        if _valor(linea, campos, "clave_tramite") != "2":
            return None
        return GRUPO_DE_TIPO.get(_valor(linea, campos, "cod_tipo"))

    todas = [linea for mes in MESES for linea in lineas_transferencias(mes)]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _hoja_mensual(
        wb,
        "V_2",
        f"Tabla 44.- V.2. Cambios de titularidad... Año {ANIO}.",
        _recuento(todas, grupo_transferencia),
        False,
    )
    wb.save(carpeta / f"transferencias_tablas_{ANIO}.xlsx")

    def grupo_baja(linea: str, campos: dict) -> str | None:
        if _valor(linea, campos, "clave_tramite") not in ("3", "4", "7"):
            return None
        return GRUPO_DE_TIPO.get(_valor(linea, campos, "cod_tipo"))

    bajas = [linea for mes in MESES for linea in lineas_bajas(mes)]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _hoja_mensual(
        wb, "V_3_Bajas", f"Tabla 49.- V.3. Bajas de vehículos... Año {ANIO}.", _recuento(bajas, grupo_baja), True
    )
    wb.save(carpeta / f"bajas_tablas_{ANIO}.xlsx")

    # Serie histórica: los años anteriores son inventados pero crecientes; el último cuadra con la tabla mensual.
    cuenta = _recuento(todas, grupo_transferencia)
    por_grupo = {g: sum(v for (_, c), v in cuenta.items() if c == g) for g in {*GRUPO_DE_TIPO.values()}}
    wb = openpyxl.Workbook()
    hoja = wb.active
    hoja.title = "cambios_Titularidad"
    hoja.append(["Tabla 14.- CAMBIOS DE TITULARIDAD. Series históricas."])
    hoja.append([])
    hoja.append(
        [
            "Año",
            "Camiones y\nfurgonetas",
            "Autobuses",
            "Turismos",
            "Motocicletas",
            "Tractores\nIndustriales",
            "Remolques y\nSemirremolques",
            "Otros\nVehículos2",
            "Total",
        ]
    )
    for anio in range(1990, ANIO):
        base = 1000 * (anio - 1989)
        fila = [base, 10, 5 * base, 100, 20, 0, 30]
        hoja.append([anio, *fila, sum(fila)])
    ultimo = [
        por_grupo["Camiones"] + por_grupo["Furgonetas"],
        0,
        por_grupo["Turismos"],
        por_grupo["Motocicletas"],
        por_grupo["Tractores\nindustriales"],
        0,
        0,
    ]
    hoja.append([ANIO, *ultimo, sum(ultimo)])
    hoja.append([])
    hoja.append(["¹ Nota al pie."])
    wb.save(carpeta / f"transferencias_series_{ANIO}.xlsx")
    return {"turismos_transferencias": por_grupo["Turismos"], "total_transferencias": sum(ultimo)}


def _documento_eurostat(indice: list[float], tasa: bool, actualizado: str) -> dict:
    periodos = [f"{2023 + (i // 12)}-{i % 12 + 1:02d}" for i in range(len(indice))]
    valores = indice
    if tasa:
        valores = [round((indice[i] / indice[i - 12] - 1) * 100, 1) if i >= 12 else None for i in range(len(indice))]
    return {
        "version": "2.0",
        "class": "dataset",
        "id": ["freq", "unit", "coicop18", "geo", "time"],
        "size": [1, 1, 1, 1, len(periodos)],
        "dimension": {"time": {"category": {"index": {p: i for i, p in enumerate(periodos)}}}},
        "value": {str(i): v for i, v in enumerate(valores) if v is not None},
        "updated": actualizado,
    }


def escribir_eurostat(raw: Path, catalogo: Path, hoy: date) -> None:
    import csv

    carpeta = raw / "eurostat" / hoy.isoformat()
    carpeta.mkdir(parents=True, exist_ok=True)
    with catalogo.open(encoding="utf-8") as f:
        series = list(csv.DictReader(f))
    for n, serie in enumerate(series):
        indice = [round(95 + n + i * 0.37, 2) for i in range(36)]  # 2023-01 a 2025-12
        documento = _documento_eurostat(indice, serie["medida"] == "tasa_anual", "2026-09-17T11:00:00+0200")
        (carpeta / f"{serie['id']}.json").write_text(json.dumps(documento), encoding="utf-8")


def construir_raw(raw: Path, catalogo_eurostat: Path, hoy: date | None = None) -> dict[str, int]:
    """Escribe toda la capa raw sintética y devuelve cifras que los tests pueden comprobar."""
    escribir_microdatos(raw)
    cifras = escribir_tablas_dgt(raw)
    escribir_eurostat(raw, catalogo_eurostat, hoy or date.today())
    return cifras
